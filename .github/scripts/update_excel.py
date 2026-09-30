"""Reads issues.json (open '[Daily Status]' issues), updates daily_status.xlsx, writes results.json."""
import csv, json, re
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

XLSX, EMP, DATE_FMT = "daily_status.xlsx", "employees.csv", "%d-%m-%Y"
MAX_BACK_DAYS, MAX_LEN = 2, 5000
HEADERS = ["#", "Date", "Completed Today", "Blockers"]
PATTERN = re.compile(r"\s*EMP_ID:\s*(\S+)\s*\nDATE:\s*(\d{2}-\d{2}-\d{4})\s*\nCOMPLETED:\n(.*)\nBLOCKERS:\n?(.*)\Z", re.S)


def style(ws):
    t = Side(style="thin", color="999999"); b = Border(left=t, right=t, top=t, bottom=t)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center"); c.border = b
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10); c.border = b
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c.column <= 2 else "left")
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"                      # never treat user text as a formula
    for col, w in zip("ABCD", (6, 14, 70, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def rows_of(ws):
    return {r[1]: (r[1], r[2] or "", r[3] or "-") for r in ws.iter_rows(min_row=2, values_only=True) if r[1]}


def clean(s):
    return "".join(ch for ch in s.replace("\r\n", "\n") if ch == "\n" or ord(ch) >= 32).strip()


def main():
    issues = json.load(open("issues.json", encoding="utf-8"))
    emps = {r["emp_id"].strip(): r["name"].strip() for r in csv.DictReader(open(EMP, encoding="utf-8"))}
    today = datetime.now(ZoneInfo("Asia/Kolkata")).date()
    wb = load_workbook(XLSX) if __import__("os").path.exists(XLSX) else Workbook()
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        wb.remove(wb["Sheet"])
    results = []
    for it in issues:
        n = it["number"]
        m = PATTERN.match((it.get("body") or "").replace("\r\n", "\n"))
        if not m:
            results.append({"number": n, "message": "Failed: could not read the entry. Please submit again from the portal."}); continue
        emp, day, done, block = m.group(1), m.group(2), clean(m.group(3)), clean(m.group(4)) or "-"
        try:
            d = datetime.strptime(day, DATE_FMT).date()
        except ValueError:
            results.append({"number": n, "message": "Failed: invalid date."}); continue
        if emp not in emps:
            results.append({"number": n, "message": f"Failed: Employee ID {emp} is not registered."}); continue
        if not (today - timedelta(days=MAX_BACK_DAYS) <= d <= today):
            results.append({"number": n, "message": f"Failed: {day} is outside the allowed window (today or up to {MAX_BACK_DAYS} days back)."}); continue
        if not done or len(done) > MAX_LEN or len(block) > MAX_LEN:
            results.append({"number": n, "message": f"Failed: 'Completed Today' must be 1-{MAX_LEN} characters."}); continue
        ws = wb[emp] if emp in wb.sheetnames else wb.create_sheet(emp)
        rows = rows_of(ws); rows[day] = (day, done, block)
        ws.delete_rows(1, ws.max_row); ws.append(HEADERS)
        for i, r in enumerate(sorted(rows.values(), key=lambda r: datetime.strptime(r[0], DATE_FMT)), 1):
            ws.append([i, *r])
        style(ws)
        results.append({"number": n, "message": f"Saved: {emps[emp]} ({emp}) - status for {day} recorded in daily_status.xlsx."})
    wb.save(XLSX)
    json.dump(results, open("results.json", "w"))
    print(json.dumps(results, indent=1))


main()
